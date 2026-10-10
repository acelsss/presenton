from copy import deepcopy
from io import BytesIO
import json

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from templates.pack.compiler import PackCompiler, source_hash, validate_pack
from templates.pack.models import TemplatePack
from templates.pack.pptx_import import extract_pack


@pytest.fixture
def source_pptx():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(5), Inches(2))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor.from_string("CCDDEE")
    text = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(4), Inches(1))
    run = text.text_frame.paragraphs[0].add_run()
    run.text = "原始标题"
    run.font.name, run.font.size = "Arial", Pt(24)
    second = deck.slides.add_slide(deck.slide_layouts[6])
    table = second.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6), Inches(3)).table
    for index, row in enumerate(table.rows):
        for col, cell in enumerate(row.cells):
            cell.text = f"{index}:{col}"
    result = BytesIO()
    deck.save(result)
    return result.getvalue()


@pytest.fixture
def pack(source_pptx):
    source, _ = extract_pack(source_pptx, pack_id="sample", version="1", filename="sample.pptx")
    return TemplatePack.model_validate(source).model_dump(mode="json")


def test_import_full_inventory_and_source_hash(pack, source_pptx):
    source, blobs = extract_pack(source_pptx, pack_id="sample", version="1", filename="sample.pptx")
    assert len(pack["layouts"]) == len(pack["pagePatterns"]) == 2
    assert len(pack["compositions"]) >= 3
    assert source["provenance"]["shapeCount"] == 3
    assert any(item["role"] == "source" and blobs[item["id"]] == source_pptx for item in source["assets"])
    assert validate_pack(pack)["checks"]["structureAndCompile"] == "pass"
    assert validate_pack(pack)["checks"]["pptxExport"] == "unrun"
    assert source_hash(json.loads(json.dumps(pack))) == source_hash(pack)


def test_template_binding_preserves_source_and_schema_rejects_unknown(pack):
    before = deepcopy(pack)
    compiler = PackCompiler(pack)
    schema = compiler.content_schema("slide_1")
    values = {key: "新标题" for key in schema["properties"]}
    compiled = compiler.compile(mode="template", ref="slide_1", content=values)
    assert "新标题" in json.dumps(compiled["ui"], ensure_ascii=False)
    assert pack == before
    assert set(values) <= compiled["nodePaths"].keys()
    with pytest.raises(Exception):
        compiler.compile(mode="template", ref="slide_1", content={"unknown": "x"})
    with pytest.raises(ValueError):
        compiler.compile(mode="template", ref="slide_1", root=pack["layouts"][0]["root"])


def test_arbitrary_composition_and_relative_vector_position(pack):
    root = {"id": "free", "type": "group", "position": {"x": 100, "y": 50}, "children": [{"id": "new_shape", "type": "vector", "position": {"x": 30, "y": 20}, "props": {"points": [{"x": 0, "y": 0}, {"x": 80, "y": 60}], "closed": False}}]}
    compiled = PackCompiler(pack).compile(mode="composition", root=root)
    native = compiled["ui"]["components"][0]["elements"][0]
    assert native["position"] == {"x": 100, "y": 50}
    assert native["children"][0]["points"] == [{"x": 30, "y": 20}, {"x": 110, "y": 80}]
    assert compiled["nodePaths"]["new_shape"].endswith("children.0")


def test_outside_canvas_visible_and_foreign_canvas_not_certified(pack):
    compiled = PackCompiler(pack).compile(mode="composition", root={"id": "outside", "type": "text", "position": {"x": 1270, "y": 710}, "size": {"width": 80, "height": 60}, "content": {"text": "Outside"}})
    assert compiled["findings"][0]["code"] == "outside_canvas"
    pack["canvas"]["height"] = 960
    assert validate_pack(pack)["checks"]["structureAndCompile"] == "fail"


def test_imported_text_inherits_recipe_and_table_keeps_tracks(pack):
    text = next(node for node in pack["layouts"][0]["root"]["children"] if node["type"] == "text")
    assert all(not run.get("font") for run in text["props"]["runs"])
    style = next(style for style in pack["elementStyles"] if style["id"] == text["styleRef"])
    style["properties"]["font"]["size"] = 42
    ui = PackCompiler(pack).compile(mode="template", ref="slide_1")["ui"]
    native = next(node for node in ui["components"][0]["elements"][0]["children"] if node["type"] == "text")
    assert native["font"]["size"] == 42
    table = next(node for node in pack["layouts"][1]["root"]["children"] if node["type"] == "table")
    assert len(table["props"]["column_widths"]) == 2
    assert len(table["props"]["row_heights"]) == 3


@pytest.mark.parametrize("corruption", ["token", "inheritance", "type", "duplicate", "binding", "unknown_prop", "composition_cycle"])
def test_semantic_validation_fails_visible(pack, corruption):
    style = pack["elementStyles"][0]
    if corruption == "token":
        style["properties"] = {"fill": {"color": "$colors.absent"}}
    elif corruption == "inheritance":
        style["extends"] = style["id"]
    elif corruption == "type":
        style["type"] = "text" if style["type"] == "vector" else "vector"
    elif corruption == "duplicate":
        pack["layouts"].append(deepcopy(pack["layouts"][0]))
    elif corruption == "binding":
        pack["layouts"][0]["bindings"]["bad"] = {"nodeId": "absent", "property": "content.text", "valueType": "string", "required": True}
    elif corruption == "unknown_prop":
        pack["layouts"][0]["root"]["children"][0]["props"]["script"] = "alert(1)"
    else:
        component = pack["compositions"][0]
        component["root"] = {"id": "cycle", "type": "instance", "compositionRef": component["id"]}
    report = validate_pack(pack)
    assert report["checks"]["structureAndCompile"] == "fail"
    assert report["findings"][-1]["severity"] == "error"


def test_instances_have_stable_noncolliding_ids(pack):
    root = {"id": "root", "type": "group", "children": [{"id": name, "type": "instance", "compositionRef": pack["compositions"][0]["id"]} for name in ("a", "b")]}
    compiler = PackCompiler(pack)
    first = compiler.compile(mode="composition", root=root)
    assert first["nodePaths"] == compiler.compile(mode="composition", root=root)["nodePaths"]
    assert any(key.startswith("a/") for key in first["nodePaths"])
    assert any(key.startswith("b/") for key in first["nodePaths"])


def test_external_image_not_accepted(pack):
    with pytest.raises(ValueError, match="pinned pack asset"):
        PackCompiler(pack).compile(mode="composition", root={"id": "external", "type": "image", "content": {"src": "https://example.com/image.png"}})
