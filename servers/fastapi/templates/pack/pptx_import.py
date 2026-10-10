"""Read-only OOXML extraction into a pack; preserve losses as diagnostics."""
from collections import Counter
from copy import deepcopy
from io import BytesIO
import hashlib
import mimetypes
from pathlib import PurePosixPath
import unicodedata
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

from templates.pack.compiler import STYLE_KEYS


MAX_UPLOAD_BYTES = 40 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 160 * 1024 * 1024
MAX_SLIDES = 100


def extract_pack(data: bytes, *, pack_id: str, version: str, filename: str):
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("PPTX exceeds 40 MiB")
    blobs, assets, diagnostics = {}, [], []
    theme_colors = {}
    with ZipFile(BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError("Expanded PPTX exceeds 160 MiB")
        if len(archive.infolist()) > 10000:
            raise ValueError("Too many ZIP entries")
        if "ppt/theme/theme1.xml" in archive.namelist():
            xml = etree.fromstring(archive.read("ppt/theme/theme1.xml"), etree.XMLParser(resolve_entities=False, no_network=True))
            for color in xml.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme/*"):
                if len(color):
                    theme_colors[etree.QName(color).localname] = "#" + (color[0].get("val") if etree.QName(color[0]).localname == "srgbClr" else color[0].get("lastClr", "000000"))
        def add_asset(blob, name, role="image"):
            digest = hashlib.sha256(blob).hexdigest()
            key = "asset_" + digest[:20]
            if key not in blobs:
                ext = PurePosixPath(name).suffix.lower()
                blobs[key] = blob
                assets.append({"id": key, "sha256": digest, "mimeType": mimetypes.guess_type(name)[0] or "application/octet-stream", "filename": digest + ext, "size": len(blob), "role": role})
            return "asset:" + key
        for item in archive.infolist():
            if item.filename.startswith("ppt/media/") and not item.is_dir():
                add_asset(archive.read(item), item.filename, "media" if item.filename.lower().endswith((".mp4", ".wmv", ".avi", ".mov", ".mp3")) else "image")
    source_ref = add_asset(data, "source.pptx", "source")
    presentation = Presentation(BytesIO(data))
    if not 0 < len(presentation.slides) <= MAX_SLIDES:
        raise ValueError("Expected 1 to 100 slides")
    scale = 1280 / presentation.slide_width
    canvas = {"width": 1280, "height": round(presentation.slide_height * scale, 4)}
    palette, fonts, sizes, radii, styles = {}, {}, {}, {}, []
    style_keys = {}
    font_usage = Counter()
    shape_count = 0

    def warn(code, slide, node_id, message):
        diagnostics.append({"severity": "warning", "code": code, "sourceSlide": slide, "nodeId": node_id, "message": message})

    def color(value, fallback="#202020"):
        try:
            if value.type == MSO_COLOR_TYPE.RGB:
                return "#" + str(value.rgb)
            if value.type == MSO_COLOR_TYPE.SCHEME:
                element = value._xClr
                return theme_colors.get(element.get("val"), fallback)
        except (AttributeError, ValueError, TypeError):
            pass
        return fallback

    def fill(fill_obj):
        if fill_obj.type == MSO_FILL_TYPE.SOLID:
            result = {"color": color(fill_obj.fore_color, "#FFFFFF")}
            alpha = fill_obj._xPr.xpath(".//a:alpha")
            if alpha:
                result["opacity"] = int(alpha[0].get("val")) / 100000
            return result
        return None

    def font(font_obj, defaults=None):
        result = dict(defaults or {})
        if font_obj.name:
            result["family"] = font_obj.name
            font_usage[font_obj.name] += 1
        if font_obj.size:
            result["size"] = round(int(font_obj.size) * scale, 3)
        if font_obj.color.type:
            result["color"] = color(font_obj.color)
        for key in ("bold", "italic", "underline"):
            value = getattr(font_obj, key)
            if value is not None:
                result[key] = bool(value)
        return result

    def tokenize(value, path=""):
        if isinstance(value, dict):
            return {key: tokenize(item, key) for key, item in value.items()}
        if isinstance(value, list):
            if path == "corner_radii":
                result = []
                for radius in value:
                    key = "radius_" + str(round(radius, 4)).replace(".", "_")
                    radii[key] = radius
                    result.append("$radii." + key)
                return result
            return [tokenize(item, path) for item in value]
        if path == "color" and isinstance(value, str) and value.startswith("#"):
            key = next((key for key, item in palette.items() if item == value), None) or f"color_{len(palette)+1:02d}"
            palette[key] = value
            return "$colors." + key
        if path == "family" and isinstance(value, str):
            key = next((key for key, item in fonts.items() if item == value), None) or f"font_{len(fonts)+1}"
            fonts[key] = value
            return "$fonts." + key
        if path == "size" and isinstance(value, (int, float)):
            key = "size_" + str(value).replace(".", "_")
            sizes[key] = value
            return "$typography." + key
        return value

    def style_for(kind, properties):
        import json
        properties = tokenize(properties)
        key = kind + json.dumps(properties, sort_keys=True)
        if key not in style_keys:
            style_id = f"{kind}_style_{len(styles)+1:03d}"
            style_keys[key] = style_id
            styles.append({"id": style_id, "name": style_id, "type": kind, "properties": properties})
        return style_keys[key]

    def native_shape(shape, slide_no, prefix):
        nonlocal shape_count
        shape_count += 1
        node_id = f"{prefix}s{shape.shape_id}"
        x, y, width, height = [round(int(v) * scale, 4) for v in (shape.left, shape.top, shape.width, shape.height)]
        base = {"id": node_id, "position": {"x": x, "y": y}, "size": {"width": max(width, .01), "height": max(height, .01)}, "source": {"slide": slide_no, "shapeId": shape.shape_id, "name": shape.name, "shapeType": str(shape.shape_type)}}
        nodes = []
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            # Explicitly refuse an inaccurate group transform until supported.
            warn("group_transform_unverified", slide_no, node_id, "Group coordinate transform is not yet certified; inspect source before reuse.")
            base.update(type="group", children=[])
            for child in shape.shapes:
                base["children"].extend(native_shape(child, slide_no, node_id + "/"))
            return [base]
        if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.MEDIA):
            try:
                image = shape.image if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else shape.part.related_part(shape._element.xpath(".//a:blip")[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"))
                blob = image.blob
                ext = image.ext if hasattr(image, "ext") else "png"
                base.update(type="image", content={"src": add_asset(blob, "poster." + ext)}, props={"fit": "fill", "rotation": shape.rotation})
                nodes.append(base)
            except (ValueError, AttributeError, IndexError, KeyError) as exc:
                warn("image_unresolved", slide_no, node_id, str(exc))
            if shape.shape_type == MSO_SHAPE_TYPE.MEDIA:
                warn("media_poster_only", slide_no, node_id, "Video preserved as an asset; native template renders its poster only.")
            elif any(getattr(shape, "crop_" + edge, 0) for edge in ("left", "right", "top", "bottom")):
                warn("image_crop_unverified", slide_no, node_id, "Original crop requires visual comparison.")
            return nodes
        if shape.has_table:
            def cell_data(cell):
                runs = []
                for p_index, paragraph in enumerate(cell.text_frame.paragraphs):
                    if p_index:
                        runs.append({"text": "\n"})
                    runs.extend({"text": run.text, "font": font(run.font, font(paragraph.font, {"family": "Noto Sans SC", "size": 18, "color": "#202020"}))} for run in paragraph.runs)
                return tokenize({"runs": runs, "color": fill(cell.fill)} if fill(cell.fill) else {"runs": runs})
            rows = [[cell_data(cell) for cell in row.cells] for row in shape.table.rows]
            base.update(type="table", props={"columns": rows[0], "rows": rows[1:], "column_widths": [int(column.width) * scale for column in shape.table.columns], "row_heights": [int(row.height) * scale for row in shape.table.rows]})
            nodes.append(base)
            if any(cell.is_merge_origin or cell.is_spanned for row in shape.table.rows for cell in row.cells):
                warn("table_merge_unverified", slide_no, node_id, "Merged table cells need source comparison.")
            warn("table_borders_defaulted", slide_no, node_id, "Cell content, fills and track sizes preserved; native border/padding defaults still require visual comparison.")
            return nodes
        properties = {}
        try:
            shape_fill = fill(shape.fill)
            if shape_fill:
                properties["fill"] = shape_fill
            if shape.fill.type not in (None, MSO_FILL_TYPE.SOLID, MSO_FILL_TYPE.BACKGROUND):
                warn("fill_unsupported", slide_no, node_id, "Non-solid fill requires visual comparison.")
        except (AttributeError, ValueError):
            pass
        try:
            if shape.line.fill.type == MSO_FILL_TYPE.SOLID:
                properties["stroke"] = {"color": color(shape.line.color), "width": round(int(shape.line.width or 9525) * scale, 3)}
        except (AttributeError, ValueError):
            pass
        if shape.shape_type == MSO_SHAPE_TYPE.FREEFORM:
            paths = shape._element.xpath(".//a:custGeom/a:pathLst/a:path")
            for index, path in enumerate(paths):
                points, closed = [], False
                pw, ph = int(path.get("w", "1")), int(path.get("h", "1"))
                for command in path:
                    tag = etree.QName(command).localname
                    if tag == "close":
                        closed = True
                    elif tag in ("moveTo", "lnTo"):
                        point = command[0]
                        points.append({"x": float(point.get("x")) / max(pw, 1) * width, "y": float(point.get("y")) / max(ph, 1) * height})
                    elif tag in ("cubicBezTo", "quadBezTo") and points:
                        controls = [points[-1], *[{"x": float(point.get("x")) / max(pw, 1) * width, "y": float(point.get("y")) / max(ph, 1) * height} for point in command]]
                        import math
                        degree = len(controls) - 1
                        for step in range(1, 17):
                            t = step / 16
                            points.append({axis: sum(math.comb(degree, j) * (1-t)**(degree-j) * t**j * control[axis] for j, control in enumerate(controls)) for axis in ("x", "y")})
                        # Keep approximation visible without dropping the path.
                        if not any(item["code"] == "path_curve_sampled" and item.get("nodeId") == node_id for item in diagnostics):
                            warn("path_curve_sampled", slide_no, node_id, "Bezier curve sampled into editable polyline points; compare with source.")
                    else:
                        warn("path_curve_unsupported", slide_no, node_id, "Curved freeform path requires visual comparison.")
                if len(points) >= 2:
                    nodes.append({**base, "id": node_id + f"_path_{index}", "type": "vector", "styleRef": style_for("vector", properties), "props": {"points": points, "closed": closed}})
        elif properties or shape.shape_type in (MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.AUTO_SHAPE):
            if not properties:
                properties["opacity"] = 0
            shape_name = str(getattr(shape, "auto_shape_type", ""))
            props = {"points": [{"x": 0, "y": 0}, {"x": width, "y": 0}, {"x": width, "y": height}, {"x": 0, "y": height}], "closed": True}
            if "OVAL" in shape_name:
                props["shape"] = "ellipse"
            elif shape.shape_type == MSO_SHAPE_TYPE.LINE:
                props.update(points=[{"x": 0, "y": 0}, {"x": width, "y": height}], closed=False)
            elif "ROUNDED_RECTANGLE" in shape_name:
                radius = min(width, height) * (shape.adjustments[0] if len(shape.adjustments) else .16667)
                properties["corner_radii"] = [radius] * 4
            elif "RECTANGLE" not in shape_name and shape_name:
                warn("shape_geometry_approximate", slide_no, node_id, f"{shape_name} approximated by its bounds.")
            props["rotation"] = shape.rotation
            nodes.append({**base, "type": "vector", "styleRef": style_for("vector", properties), "props": props})
        if shape.has_text_frame and shape.text:
            frame = shape.text_frame
            runs, default_font = [], None
            for p_index, paragraph in enumerate(frame.paragraphs):
                if p_index:
                    runs.append({"text": "\n"})
                for run in paragraph.runs:
                    run_font = font(run.font, font(paragraph.font, {"family": "Noto Sans SC", "size": 24, "color": "#202020"}))
                    if isinstance(paragraph.line_spacing, float):
                        run_font["line_height"] = paragraph.line_spacing
                    default_font = default_font or run_font
                    runs.append({"text": run.text, "font": tokenize(run_font)})
            first = frame.paragraphs[0]
            shared_font = tokenize(default_font or {})
            for run in runs:
                if "font" in run:
                    run["font"] = {key: value for key, value in run["font"].items() if shared_font.get(key) != value}
            alignment = str(first.alignment).split(" ")[0].lower() if first.alignment else "left"
            if alignment not in {"left", "center", "right", "justify"}:
                alignment = "left"
            text_node = deepcopy(base)
            text_node["id"] += "_text"
            text_node["position"] = {"x": x + int(frame.margin_left) * scale, "y": y + int(frame.margin_top) * scale}
            text_node["size"] = {"width": max(.01, width - int(frame.margin_left + frame.margin_right) * scale), "height": max(.01, height - int(frame.margin_top + frame.margin_bottom) * scale)}
            text_node["source"]["textFrame"] = {"wordWrap": frame.word_wrap, "position": deepcopy(text_node["position"]), "size": deepcopy(text_node["size"])}
            # The native editor deliberately wraps words. Source no-wrap boxes
            # can be narrower than their visible text. Expand those bounds
            # conservatively and record the adaptation, without shrinking type.
            if frame.word_wrap is False:
                em = (default_font or {}).get("size", 24)
                estimate = max(sum(1 if unicodedata.east_asian_width(char) in ("W", "F") else .7 for char in line) for line in shape.text.split("\n")) * em + 8
                if estimate > text_node["size"]["width"]:
                    old_width = text_node["size"]["width"]
                    new_width = min(estimate, canvas["width"] - text_node["position"]["x"])
                    delta = max(0, new_width - old_width)
                    if alignment == "center":
                        text_node["position"]["x"] -= delta / 2
                    elif alignment == "right":
                        text_node["position"]["x"] -= delta
                    text_node["size"]["width"] = max(old_width, new_width)
                    warn("text_wrap_adapted", slide_no, text_node["id"], "Source no-wrap text bounds expanded for native word wrapping; verify font metrics and neighbouring elements.")
            vertical = str(frame.vertical_anchor).split(" ")[0].lower() if frame.vertical_anchor else "top"
            text_node.update(type="text", styleRef=style_for("text", {"font": default_font or {}, "alignment": {"horizontal": alignment, "vertical": vertical if vertical in ("top", "middle", "bottom") else "top"}}), content={"text": shape.text}, props={"runs": runs, "rotation": shape.rotation})
            nodes.append(text_node)
        if not nodes:
            warn("shape_not_rendered", slide_no, node_id, f"No native representation for {shape.shape_type}; source PPTX retained.")
        return nodes

    layouts, compositions, patterns = [], [], []
    for slide_no, slide in enumerate(presentation.slides, 1):
        elements = []
        background = fill(slide.background.fill) or {"color": "#FFFFFF"}
        elements.append({"id": f"p{slide_no}_background", "type": "vector", "position": {"x": 0, "y": 0}, "styleRef": style_for("vector", {"fill": background}), "props": {"points": [{"x": 0, "y": 0}, {"x": 1280, "y": 0}, {"x": 1280, "y": canvas["height"]}, {"x": 0, "y": canvas["height"]}], "closed": True}})
        for shape in slide.shapes:
            elements.extend(native_shape(shape, slide_no, f"p{slide_no}_"))
        title = next((shape.text for shape in slide.shapes if shape.has_text_frame and shape.text), f"第 {slide_no} 页")
        root = {"id": f"page_{slide_no}", "type": "group", "size": canvas, "children": elements}
        bindings = {}
        for node in elements:
            if node["type"] == "text":
                bindings[node["id"]] = {"nodeId": node["id"], "property": "content.text", "valueType": "string", "required": True}
            elif node["type"] == "table":
                bindings[node["id"]] = {"nodeId": node["id"], "property": "props.rows", "valueType": "array", "required": True}
        layouts.append({"id": f"slide_{slide_no}", "name": title[:100], "root": root, "bindings": bindings, "sourceSlide": slide_no})
        compositions.append({"id": f"page_composition_{slide_no}", "name": title[:100], "root": deepcopy(root)})
        patterns.append({"id": f"pattern_{slide_no}", "name": title[:100], "compositionRef": f"page_composition_{slide_no}"})
        # Geometric candidates are labelled as such, never certified semantics.
        for node in elements:
            if node["type"] != "vector" or not node.get("size"):
                continue
            box = node["size"]
            if not (120 <= box["width"] < 1100 and 65 <= box["height"] < 600):
                continue
            nx, ny = node["position"]["x"], node["position"]["y"]
            texts = [item for item in elements if item["type"] == "text" and nx <= item["position"]["x"] < nx + box["width"] and ny <= item["position"]["y"] < ny + box["height"]]
            if not texts:
                continue
            children = deepcopy([node, *texts])
            for child in children:
                child["position"]["x"] -= nx
                child["position"]["y"] -= ny
            compositions.append({"id": "cluster_" + node["id"], "name": f"第 {slide_no} 页 · 空间组合候选 {node['source']['shapeId']}", "root": {"id": "cluster_" + node["id"], "type": "group", "size": box, "children": children}})
    pack = {"schemaVersion": "1.0", "id": pack_id, "version": version, "name": PurePosixPath(filename.replace("\\", "/")).stem, "supportedModes": ["template", "composition"], "canvas": canvas, "defaultThemeRef": "source_theme", "themes": [{"id": "source_theme", "name": "原稿主题", "tokens": {"colors": palette, "fonts": fonts, "typography": sizes, "spacing": {}, "radii": radii}}], "elementStyles": styles, "compositions": compositions, "pagePatterns": patterns, "layouts": layouts, "assets": assets, "provenance": {"filename": PurePosixPath(filename.replace("\\", "/")).name, "sourceSha256": hashlib.sha256(data).hexdigest(), "sourceAssetRef": source_ref, "slideCount": len(layouts), "shapeCount": shape_count, "fonts": dict(font_usage), "extractor": "python-pptx-ooxml-v1", "note": "Absolute geometry extracted; geometric clusters require human review. Spacing roles and semantic diagrams are not inferred."}, "diagnostics": diagnostics}
    return pack, blobs
